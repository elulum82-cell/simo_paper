"""
SiMo SupeR PapeR (v3)

Alur:
  1. Bahan   : unggah dokumen -> "Baca dokumen" menjalankan skill paper-generator dan menyimpan draf
               (inventaris, klaim inti, kerangka, catatan perbaikan).
  2. Rancang : tiap catatan bisa ditangani dengan (a) mengunggah dokumen lain atau (b) mencari data
               di sumber terpercaya.
  3. Tulis   : "Tulis semua bagian" menyusun paper lengkap dari draf + data tambahan; judul/abstrak/kata kunci.
  4. Review  : review hasil tulis -> perbaiki paper sesuai review -> daftar pustaka -> ekspor.

Jalankan:
    pip install -r requirements.txt
    export ANTHROPIC_API_KEY="sk-ant-..."   # atau isi di sidebar
    streamlit run app.py
Skill dibaca dari folder ./skill (bawaan) atau dari berkas .skill yang diunggah di sidebar.
"""
import html
import io
import json
import os
import re
import zipfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
import streamlit as st  # noqa: E402
from anthropic import Anthropic  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from pypdf import PdfReader  # noqa: E402

APP_DIR = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.join(APP_DIR, "skill")

# ----------------------------------------------------------------------------
# Konfigurasi
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "claude-sonnet-5-5"
FILE_TYPES = ["pdf", "docx", "xlsx", "xls", "csv", "txt", "md"]

PAPER_TYPES = {
    "Penelitian empiris": ["Pendahuluan", "Metode", "Hasil", "Pembahasan", "Kesimpulan"],
    "Studi kasus / kajian kebijakan-teknis": [
        "Pendahuluan", "Latar dan Profil Kasus", "Metode", "Temuan", "Pembahasan", "Kesimpulan dan Rekomendasi",
    ],
    "Pemodelan / simulasi": [
        "Pendahuluan", "Struktur Model dan Metode", "Kalibrasi dan Validasi", "Hasil Simulasi dan Skenario",
        "Pembahasan", "Kesimpulan",
    ],
    "Tinjauan literatur": ["Pendahuluan", "Metode Review", "Hasil Sintesis", "Pembahasan", "Kesimpulan"],
    "Teknis / rekayasa": ["Pendahuluan", "Karya Terkait", "Metode yang Diusulkan", "Evaluasi", "Pembahasan", "Kesimpulan"],
}

SECTION_GUIDE = {
    "Pendahuluan": (
        "Pola corong: konteks luas -> masalah -> yang sudah dilakukan peneliti lain (dengan sitasi) -> GAP -> "
        "tujuan dan kontribusi -> peta isi paper. Gap harus didukung sitasi."
    ),
    "Metode": (
        "Cukup rinci untuk replikasi: objek/lokasi, sumber dan periode data, variabel, rumus/model, asumsi, "
        "alat, validasi. Jelaskan alasan memilih metode."
    ),
    "Hasil": (
        "Sajikan temuan objektif tanpa tafsir panjang. Gunakan tabel Markdown bila perlu; beri nomor dan "
        "keterangan mandiri dan rujuk di teks. Sertakan satuan dan ketidakpastian."
    ),
    "Pembahasan": (
        "Tafsirkan hasil, bandingkan dengan literatur (sesuai/bertentangan dan alasannya), implikasi praktis/"
        "kebijakan, KETERBATASAN yang jujur, dan riset lanjutan."
    ),
    "Kesimpulan": "Jawab pertanyaan penelitian satu per satu, singkat, tanpa data atau argumen baru.",
}

BASE_SYSTEM = """Anda adalah penulis ilmiah senior yang menyusun manuskrip jurnal dari bahan mentah pengguna.

PRINSIP WAJIB:
1. Jangan mengarang. Angka, hasil, nama, dan sitasi hanya boleh berasal dari bahan pengguna, dokumen tambahan, atau sumber yang benar-benar Anda temukan lewat pencarian. Bila data tidak ada, tulis [DATA DIBUTUHKAN: ...] dan jangan menebak.
2. Setiap klaim punya sumber. Klaim dari bahan pengguna dirujuk (mis. [DOK-1]); klaim dari literatur diberi sitasi (Penulis, Tahun). Sitasi yang tidak bisa diverifikasi ditandai [VERIFIKASI].
3. Ikuti bahasa naskah yang diminta; jangan mencampur bahasa.
4. Jujur soal keterbatasan data dan metode.
5. Gaya: kalimat aktif dan ringkas, satu ide per paragraf, istilah konsisten, singkatan dijelaskan saat pertama muncul, hindari klaim berlebihan.
6. Jangan menyalin teks sumber; parafrasekan dan sitasi.
7. Gunakan Markdown (tabel dengan sintaks pipa bila perlu). Jangan menambahkan basa-basi di awal atau akhir."""

APP_NOTE = (
    "\n\nCATATAN PENERAPAN: Anda dipanggil oleh aplikasi, bukan percakapan. Jangan bertanya balik dan jangan "
    "menulis basa-basi. Abaikan bagian skill tentang membuat file, memakai skill lain, atau menyerahkan file; "
    "kerjakan tugas pada pesan pengguna dan patuhi format keluaran yang diminta."
)


# ----------------------------------------------------------------------------
# Skill
# ----------------------------------------------------------------------------
def _strip_frontmatter(t: str) -> str:
    return re.sub(r"^---\n.*?\n---\n", "", t, count=1, flags=re.S)


def load_skill(uploaded=None):
    """Kembalikan (teks_skill, keterangan). Dari berkas .skill yang diunggah, atau folder ./skill."""
    docs = []
    try:
        if uploaded is not None:
            with zipfile.ZipFile(io.BytesIO(uploaded.getvalue())) as z:
                for n in sorted(z.namelist()):
                    if n.endswith(".md"):
                        docs.append((n, z.read(n).decode("utf-8", "ignore")))
        else:
            for root, _, fs in os.walk(SKILL_DIR):
                for fn in sorted(fs):
                    if fn.endswith(".md"):
                        p = os.path.join(root, fn)
                        with open(p, encoding="utf-8") as fh:
                            docs.append((os.path.relpath(p, SKILL_DIR), fh.read()))
    except Exception as e:  # noqa: BLE001
        return "", f"gagal memuat skill: {e}"
    main = [d for d in docs if d[0].endswith("SKILL.md")]
    if not main:
        return "", "SKILL.md tidak ditemukan"
    text = _strip_frontmatter(main[0][1])
    for p, t in docs:
        if not p.endswith("SKILL.md"):
            text += f"\n\n--- Referensi skill: {p} ---\n{t}"
    return text, f"{len(text):,} karakter"


# ----------------------------------------------------------------------------
# Pembaca dokumen
# ----------------------------------------------------------------------------
def extract_text(f) -> str:
    name = f.name.lower()
    data = f.getvalue()
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(f"[hal {i + 1}] " + (p.extract_text() or "") for i, p in enumerate(reader.pages))
        if name.endswith(".docx"):
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for t in doc.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)
        if name.endswith((".xlsx", ".xls")):
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
            return "\n\n".join(_df_summary(df, f"Sheet '{n}'") for n, df in sheets.items())
        if name.endswith(".csv"):
            return _df_summary(pd.read_csv(io.BytesIO(data)), "CSV")
        return data.decode("utf-8", errors="ignore")
    except Exception as e:  # noqa: BLE001
        return f"[GAGAL MEMBACA {f.name}: {e}]"


def _df_summary(df: pd.DataFrame, label: str) -> str:
    out = [f"{label}: {df.shape[0]} baris x {df.shape[1]} kolom", "Kolom: " + ", ".join(map(str, df.columns))]
    out.append(df.head(100).to_csv(index=False))
    num = df.select_dtypes("number")
    if not num.empty:
        out.append("Statistik deskriptif:\n" + num.describe().round(4).to_string())
    return "\n".join(out)


def extract_tables(files, prefix: str = "DOK") -> dict:
    """Ambil tabel (CSV/XLSX) sebagai DataFrame untuk dibuat grafik."""
    out = {}
    for i, f in enumerate(files, 1):
        n = f.name.lower()
        try:
            if n.endswith(".csv"):
                out[f"[{prefix}-{i}] {f.name}"] = pd.read_csv(io.BytesIO(f.getvalue()))
            elif n.endswith((".xlsx", ".xls")):
                for sn, df in pd.read_excel(io.BytesIO(f.getvalue()), sheet_name=None).items():
                    out[f"[{prefix}-{i}] {f.name} / {sn}"] = df
        except Exception:  # noqa: BLE001
            pass
    return out


def make_chart(df: pd.DataFrame, kind: str, x, ys: list) -> bytes:
    d = df[[x] + ys].dropna()
    if kind == "Batang" and len(d) > 40:
        d = d.head(40)
    fig, ax = plt.subplots(figsize=(7, 4), dpi=200)
    if kind == "Garis":
        for y in ys:
            ax.plot(d[x], d[y], marker="o", markersize=3, label=str(y))
    elif kind == "Batang":
        d.set_index(x)[ys].plot(kind="bar", ax=ax)
        ax.tick_params(axis="x", rotation=45)
    else:
        for y in ys:
            ax.scatter(d[x], d[y], s=14, label=str(y))
    ax.set_xlabel(str(x))
    ax.set_ylabel(", ".join(map(str, ys)))
    ax.grid(alpha=0.3)
    if len(ys) > 1:
        ax.legend()
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def build_context(files, per_doc_cap: int, total_cap: int, prefix: str = "DOK") -> str:
    chunks, used = [], 0
    for i, f in enumerate(files, 1):
        text = extract_text(f)
        if len(text) > per_doc_cap:
            text = text[:per_doc_cap] + "\n[...dipotong karena terlalu panjang...]"
        chunk = f"=== [{prefix}-{i}] {f.name} ===\n{text}\n"
        if used + len(chunk) > total_cap:
            chunk = chunk[: max(0, total_cap - used)] + "\n[...batas total konteks tercapai...]"
        chunks.append(chunk)
        used += len(chunk)
        if used >= total_cap:
            break
    return "\n".join(chunks)


# ----------------------------------------------------------------------------
# Pemanggilan model
# ----------------------------------------------------------------------------
def call_llm(user: str, system_extra: str = "", max_tokens: int = 6000, web: bool = False) -> str:
    s = st.session_state
    skill_part = f"\n\n=== SKILL: paper-generator (ikuti alur dan prinsipnya) ===\n{s.skill}" if s.skill else ""
    system = [{"type": "text", "text": BASE_SYSTEM + skill_part + APP_NOTE, "cache_control": {"type": "ephemeral"}}]
    if system_extra:
        system.append({"type": "text", "text": system_extra})
    kwargs = dict(model=s.model, max_tokens=max_tokens, system=system, messages=[{"role": "user", "content": user}])
    if web:
        kwargs["tools"] = [{"type": "web_search_20250305", "name": "web_search", "max_uses": 8}]
    try:
        resp = Anthropic(api_key=s.api_key).messages.create(**kwargs)
    except Exception as e:  # noqa: BLE001
        st.error(f"Gagal memanggil model: {e}")
        return ""
    return "".join(b.text for b in resp.content if getattr(b, "type", "") == "text").strip()


def meta_block() -> str:
    s = st.session_state
    return (
        f"Bahasa naskah: {s.lang}\nJenis paper: {s.ptype}\n"
        f"Target publikasi: {s.target or '(belum ditentukan)'}\n"
        f"Gaya sitasi: {s.cite} (format sitasi di dalam teks HARUS sesuai gaya ini)"
    )


def source_ctx() -> str:
    s = st.session_state
    return s.ctx + ("\n\n=== DOKUMEN TAMBAHAN ===\n" + s.extra_ctx[:150_000] if s.extra_ctx else "")


def additions_block() -> str:
    s = st.session_state
    done, open_ = [], []
    for n in s.notes:
        fx = s.fixes.get(n["id"])
        if fx and fx.get("hasil", "").strip():
            done.append(f"- Catatan #{n['id']} ({n['judul']}) — ditangani lewat {fx['sumber']}:\n{fx['hasil']}")
        else:
            open_.append(f"- #{n['id']} [{n['kategori']}] {n['judul']}")
    return (
        "DATA TAMBAHAN DARI TAB RANCANG:\n" + ("\n".join(done) or "(tidak ada)")
        + "\n\nCATATAN YANG MASIH TERBUKA (jangan ditutup dengan karangan; tandai [DATA DIBUTUHKAN] atau tulis "
        "sebagai keterbatasan):\n" + ("\n".join(open_) or "(tidak ada)")
    )


# ----------------------------------------------------------------------------
# Parsing keluaran analisis
# ----------------------------------------------------------------------------
def grab(text: str, tag: str, nxt: str = "") -> str:
    pat = rf"==={tag}===\s*(.*?)\s*(?:==={nxt}===|\Z)" if nxt else rf"==={tag}===\s*(.*)\Z"
    m = re.search(pat, text, re.S)
    return m.group(1).strip() if m else ""


NOTE_RE = re.compile(r"^[\s\-*\d.)]*\[(.+?)\]\s*\|\|\s*(.+?)\s*\|\|\s*(.+?)\s*\|\|\s*(.+)$")


def parse_notes(block: str) -> list:
    notes = []
    for line in block.splitlines():
        m = NOTE_RE.match(line.strip())
        if m:
            notes.append({
                "id": len(notes) + 1, "kategori": m.group(1).strip(), "judul": m.group(2).strip(),
                "detail": m.group(3).strip(), "saran": m.group(4).strip(),
            })
    return notes


# ----------------------------------------------------------------------------
# Markdown -> DOCX sederhana
# ----------------------------------------------------------------------------
def _add_runs(par, text: str):
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part:
            par.add_run(part)


def md_to_docx(md: str, figures=None) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        mg = re.match(r"^\[\[GAMBAR:(\d+)\]\]$", line.strip())
        if mg:
            idx = int(mg.group(1)) - 1
            if figures and 0 <= idx < len(figures):
                doc.add_picture(io.BytesIO(figures[idx]["png"]), width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.add_run(f"Gambar {idx + 1}. {figures[idx]['caption']}").italic = True
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncol)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(ncol):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            if level == 1:
                h = doc.add_heading(m.group(2), 0)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(m.group(2), level - 1)
        elif re.match(r"^\s*[-*]\s+", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(p, line)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def assemble_manuscript() -> str:
    s = st.session_state
    parts = [f"# {s.meta.get('judul', 'Judul Paper')}", "", "## Abstrak", s.meta.get("abstrak", ""), ""]
    if s.meta.get("kata_kunci"):
        parts += [f"**Kata kunci:** {s.meta['kata_kunci']}", ""]
    result_name = next((n for n in s.section_names if n.split()[0] in ("Hasil", "Temuan", "Evaluasi")), None)
    for name in s.section_names:
        if s.sections.get(name):
            parts += [f"## {name}", s.sections[name], ""]
            if name == result_name:
                for i in range(1, len(s.figures) + 1):
                    parts += [f"[[GAMBAR:{i}]]", ""]
    if s.refs:
        parts += ["## Daftar Pustaka", s.refs, ""]
    return "\n".join(parts)


def figs_to_text(md: str) -> str:
    def _sub(m):
        i = int(m.group(1)) - 1
        cap = st.session_state.figures[i]["caption"] if i < len(st.session_state.figures) else ""
        return f"*Gambar {i + 1}. {cap}*"
    return re.sub(r"\[\[GAMBAR:(\d+)\]\]", _sub, md)


# ----------------------------------------------------------------------------
# State
# ----------------------------------------------------------------------------
def init_state():
    d = {
        "api_key": os.environ.get("ANTHROPIC_API_KEY", ""),
        "model": DEFAULT_MODEL,
        "lang": "Indonesia",
        "ptype": list(PAPER_TYPES)[0],
        "target": "",
        "cite": "APA 7",
        "skill": "",
        "ctx": "",
        "extra_ctx": "",
        "inv": "",
        "claim": "",
        "outline": "",
        "notes": [],
        "fixes": {},
        "lit": "",
        "sections": {},
        "section_names": [],
        "meta": {},
        "refs": "",
        "qa": "",
        "ver": {},
        "tables": {},
        "figures": [],
        "backup": None,
    }
    for k, v in d.items():
        st.session_state.setdefault(k, v)


def bump(key: str):
    st.session_state.ver[key] = st.session_state.ver.get(key, 0) + 1


def editor(label: str, key: str, height: int = 300):
    """Text area yang bisa diedit; versi dinaikkan saat isi diganti oleh model."""
    ver = st.session_state.ver.get(key, 0)
    st.session_state[key] = st.text_area(label, value=st.session_state[key], height=height, key=f"ta_{key}_{ver}")


def reset_downstream():
    s = st.session_state
    s.update(
        notes=[], fixes={}, extra_ctx="", lit="", sections={}, meta={}, refs="", qa="", tables={}, figures=[], backup=None
    )
    for k in ("inv", "claim", "outline", "lit", "refs", "qa"):
        bump(k)
    for n in PAPER_TYPES_ALL_SECTIONS:
        bump(f"sec_{n}")


PAPER_TYPES_ALL_SECTIONS = sorted({n for v in PAPER_TYPES.values() for n in v})


def need_ready(need_text: bool = False) -> bool:
    s = st.session_state
    if not s.api_key:
        st.error("Isi API key di sidebar terlebih dahulu.")
        return False
    if not s.ctx:
        st.error("Unggah dokumen dan klik 'Baca dokumen' di tab Bahan.")
        return False
    if need_text and not any(s.sections.values()):
        st.error("Belum ada naskah. Jalankan 'Tulis semua bagian' di tab Tulis.")
        return False
    return True


# ----------------------------------------------------------------------------
# Operasi utama
# ----------------------------------------------------------------------------
def analyze_documents(files, per_doc: int, total_cap: int) -> bool:
    """Tab 1: baca dokumen lalu jalankan skill (Tahap 1-3) + identifikasi catatan perbaikan."""
    s = st.session_state
    with st.spinner("Membaca dokumen..."):
        ctx = build_context(files, per_doc, total_cap)
        tables = extract_tables(files)
    with st.spinner("Menganalisis bahan dengan skill paper-generator..."):
        out = call_llm(
            f"{meta_block()}\n\nBAHAN:\n{ctx}\n\n"
            "TUGAS: jalankan Tahap 1-3 skill (inventaris bahan, klaim inti, kerangka) lalu identifikasi catatan "
            "perbaikan yang menghambat paper ini diterima jurnal bereputasi.\n"
            "1) INVENTARIS: tabel Markdown (Bahan | Isi utama | Dipakai untuk bagian), lalu daftar singkat: "
            "pertanyaan yang dijawab, metode, hasil kunci (angka/tabel), dan celah data.\n"
            "2) KLAIM_INTI: 4 butir — Masalah/gap, Tujuan/pertanyaan penelitian (satu kalimat), Kontribusi/kebaruan, "
            "Temuan utama. Bila kebaruan lemah, katakan terus terang.\n"
            f"3) KERANGKA: untuk bagian {', '.join(s.section_names)}: 1-2 baris isi dan bahan sumber (DOK-n) per bagian.\n"
            "4) CATATAN: maksimal 10 catatan urut prioritas — celah data, kebaruan lemah, metode kurang rinci, "
            "kontradiksi antar dokumen/bagian, kebutuhan literatur pembanding, validasi/uji yang belum ada, "
            "etika/izin data. SATU BARIS per catatan, format persis:\n"
            "[Kategori] || Judul singkat || Detail masalah (apa yang kurang, mengapa penting, rujuk DOK-n) || "
            "Saran penyelesaian (data/dokumen apa yang perlu diunggah atau dicari)\n\n"
            "FORMAT KELUARAN PERSIS, tanpa teks lain:\n"
            "===INVENTARIS===\n...\n===KLAIM_INTI===\n...\n===KERANGKA===\n...\n===CATATAN===\n...",
            max_tokens=8000,
        )
    if not out:
        return False
    reset_downstream()
    s.ctx, s.tables = ctx, tables
    s.inv = grab(out, "INVENTARIS", "KLAIM_INTI")
    s.claim = grab(out, "KLAIM_INTI", "KERANGKA")
    s.outline = grab(out, "KERANGKA", "CATATAN")
    s.notes = parse_notes(grab(out, "CATATAN"))
    if not (s.inv or s.claim or s.outline):
        s.inv = out  # format tak terbaca: simpan mentah agar tidak hilang
        st.warning("Format keluaran model tidak terbaca penuh; hasil mentah disimpan di Inventaris.")
    return True


def handle_upload_fix(note: dict, ups, per_doc: int):
    s = st.session_state
    nid = note["id"]
    with st.spinner("Membaca dokumen tambahan..."):
        prefix = f"TAMBAHAN-{nid}"
        text = build_context(ups, per_doc, 80_000, prefix=prefix)
        s.tables.update(extract_tables(ups, prefix))
    with st.spinner("Mengekstrak informasi yang menjawab catatan..."):
        digest = call_llm(
            f"{meta_block()}\n\nCATATAN #{nid} [{note['kategori']}] {note['judul']}\nDetail: {note['detail']}\n"
            f"Saran: {note['saran']}\n\nKLAIM INTI:\n{s.claim}\n\nDOKUMEN TAMBAHAN:\n{text}\n\n"
            "TUGAS: dari dokumen tambahan saja, ekstrak informasi/data yang menjawab catatan ini (angka, metode, "
            "parafrase temuan, rujukan [TAMBAHAN-n-m]). Jelaskan singkat bagaimana ini menutup catatan, dan "
            "sebutkan apa yang MASIH kurang. Jangan memakai informasi di luar dokumen tambahan.",
            max_tokens=3000,
        )
    if digest:
        s.extra_ctx += ("\n" if s.extra_ctx else "") + text
        _store_fix(nid, "dokumen tambahan: " + ", ".join(f.name for f in ups), digest)
        bump(f"up_{nid}")


def handle_web_fix(note: dict):
    s = st.session_state
    with st.spinner("Mencari di sumber terpercaya (bisa beberapa menit)..."):
        res = call_llm(
            f"{meta_block()}\n\nCATATAN #{note['id']} [{note['kategori']}] {note['judul']}\nDetail: {note['detail']}\n"
            f"Saran: {note['saran']}\n\nKLAIM INTI:\n{s.claim}\n\nRINGKASAN BAHAN:\n{s.inv}\n\n"
            "TUGAS: cari data/informasi yang menjawab catatan ini hanya dari sumber terpercaya (jurnal bereputasi, "
            "lembaga pemerintah/statistik resmi, organisasi internasional, laporan teknis lembaga resmi). Untuk tiap "
            "temuan tulis: ringkasan data/angka, sumber (penulis/lembaga, tahun, judul, URL/DOI), dan tingkat "
            "keyakinan. Hanya cantumkan yang benar-benar ditemukan; tandai [VERIFIKASI] bila tidak yakin; bila tidak "
            "ditemukan, katakan terus terang. Jangan mengarang.",
            max_tokens=3500,
            web=True,
        )
    if res:
        _store_fix(note["id"], "pencarian web sumber terpercaya", res)


def _store_fix(nid: int, sumber: str, hasil: str):
    fx = st.session_state.fixes.get(nid)
    if fx and fx.get("hasil", "").strip():
        fx["sumber"] += " + " + sumber
        fx["hasil"] += "\n\n---\n" + hasil
    else:
        st.session_state.fixes[nid] = {"sumber": sumber, "hasil": hasil}
    bump(f"fx_{nid}")


def search_literature():
    s = st.session_state
    with st.spinner("Mencari literatur (bisa beberapa menit)..."):
        res = call_llm(
            f"{meta_block()}\n\nKLAIM INTI:\n{s.claim}\n\n{additions_block()}\n\n"
            "Cari 15-30 referensi relevan (prioritaskan jurnal bereputasi 5-10 tahun terakhir + karya klasik). "
            "Untuk tiap referensi tulis: Penulis (Tahun). Judul. Venue. DOI/URL — satu kalimat relevansinya. "
            "Hanya cantumkan yang benar-benar ditemukan; beri [VERIFIKASI] bila tidak yakin.",
            max_tokens=5000,
            web=True,
        )
    if res:
        s.lit = res
        bump("lit")


def write_section(name: str):
    s = st.session_state
    guide = SECTION_GUIDE.get(name, "Tulis sesuai fungsi bagian ini dalam paper, runtut dan berbasis bukti.")
    figs = "\n".join(f"Gambar {i}: {f['caption']}" for i, f in enumerate(s.figures, 1)) or "(tidak ada)"
    prev = "\n\n".join(f"## {n}\n{t}" for n, t in s.sections.items() if t and n != name)[-12000:]
    text = call_llm(
        f"{meta_block()}\n\nKLAIM INTI:\n{s.claim}\n\nKERANGKA:\n{s.outline}\n\n"
        f"LITERATUR:\n{s.lit or '(tidak ada; jangan mengarang sitasi)'}\n\n{additions_block()}\n\n"
        f"GAMBAR TERSEDIA (rujuk di teks bila relevan; jangan mengarang gambar lain):\n{figs}\n\n"
        f"BAHAN:\n{source_ctx()}\n\nBAGIAN SEBELUMNYA (untuk konsistensi):\n{prev or '(belum ada)'}\n\n"
        f"TUGAS: Tulis HANYA isi bagian '{name}' (tanpa heading utama). Panduan: {guide} "
        "Gunakan sub-heading level 3 (###) bila perlu.",
        max_tokens=6000,
    )
    if text:
        s.sections[name] = text
        bump(f"sec_{name}")


def make_meta():
    s = st.session_state
    body = "\n\n".join(f"## {n}\n{s.sections.get(n, '')}" for n in s.section_names)
    with st.spinner("Menyusun judul dan abstrak..."):
        out = call_llm(
            f"{meta_block()}\n\nKLAIM INTI:\n{s.claim}\n\nISI NASKAH:\n{body}\n\n"
            "Buat: judul spesifik (objek, pendekatan, konteks); abstrak satu paragraf 150-250 kata "
            "(latar singkat, tujuan, metode, hasil kunci dengan angka, kesimpulan) tanpa sitasi; "
            "4-6 kata kunci. Angka harus sama persis dengan naskah. Format keluaran persis:\n"
            "JUDUL: ...\nABSTRAK: ...\nKATA KUNCI: ...",
            max_tokens=1500,
        )
    if not out:
        return
    m = re.search(r"JUDUL:\s*(.*?)\s*ABSTRAK:\s*(.*?)\s*KATA KUNCI:\s*(.*)", out, re.S)
    if m:
        s.meta = {"judul": m.group(1).strip(), "abstrak": m.group(2).strip(), "kata_kunci": m.group(3).strip()}
    else:
        s.meta = {"judul": "Judul Paper", "abstrak": out, "kata_kunci": ""}


def run_review():
    s = st.session_state
    with st.spinner("Mereview naskah..."):
        out = call_llm(
            f"{meta_block()}\n\nNASKAH:\n{figs_to_text(assemble_manuscript())}\n\nKLAIM INTI:\n{s.claim}\n\n"
            f"{additions_block()}\n\nBAHAN SUMBER:\n{source_ctx()}\n\n"
            "TUGAS: review naskah seperti reviewer jurnal bereputasi yang teliti, memakai checklist skill. Periksa: "
            "(1) tujuan-metode-hasil-kesimpulan saling menjawab; (2) angka di abstrak/kesimpulan cocok dengan Hasil dan "
            "BAHAN SUMBER (sebutkan ketidakcocokan); (3) tabel/gambar dirujuk; (4) sitasi di teks vs daftar pustaka; "
            "(5) klaim tanpa dukungan; (6) penanda [DATA DIBUTUHKAN] dan [VERIFIKASI] yang tersisa; (7) etika (izin data, "
            "persetujuan etik, pengungkapan penggunaan AI); (8) kejelasan kebaruan dan keterbatasan.\n"
            "Format keluaran:\n## Ringkasan penilaian\n## Temuan (daftar bernomor: [Kritis/Penting/Minor] bagian "
            "terkait — masalah — perbaikan yang diminta)\n## Risiko penolakan reviewer (3-5)\n## Kekuatan paper\n"
            "## Tindakan untuk penulis (data/verifikasi yang tidak bisa diperbaiki otomatis)",
            max_tokens=4000,
        )
    if out:
        s.qa = out
        bump("qa")


def revise_all():
    s = st.session_state
    s.backup = {"sections": dict(s.sections), "meta": dict(s.meta), "refs": s.refs}
    names = [n for n in s.section_names if s.sections.get(n)]
    bar = st.progress(0.0)
    for idx, name in enumerate(names):
        others = "\n\n".join(f"## {n}\n{t}" for n, t in s.sections.items() if t and n != name)[-12000:]
        with st.spinner(f"Memperbaiki {name}..."):
            new = call_llm(
                f"{meta_block()}\n\nHASIL REVIEW:\n{s.qa}\n\nKLAIM INTI:\n{s.claim}\n\n"
                f"LITERATUR:\n{s.lit or '(tidak ada)'}\n\n{additions_block()}\n\nBAHAN:\n{source_ctx()}\n\n"
                f"BAGIAN LAIN (konteks):\n{others or '(tidak ada)'}\n\nBAGIAN YANG DIPERBAIKI: '{name}'\n"
                f"ISI SAAT INI:\n{s.sections[name]}\n\n"
                "TUGAS: perbaiki HANYA bagian ini sesuai temuan review yang relevan. Pertahankan yang sudah baik; "
                "jangan menambah angka atau sitasi yang tidak ada di bahan, data tambahan, atau literatur; bila data "
                "kurang, tandai [DATA DIBUTUHKAN]. Keluarkan hanya isi bagian yang telah diperbaiki.",
                max_tokens=6000,
            )
        if new:
            s.sections[name] = new
            bump(f"sec_{name}")
        bar.progress((idx + 1) / len(names))
    if s.meta:
        make_meta()
    s.refs = ""  # daftar pustaka perlu disusun ulang dari naskah yang baru
    bump("refs")


# ----------------------------------------------------------------------------
# Tampilan (CSS + komponen kecil) — hanya visual, tidak mengubah logika
# ----------------------------------------------------------------------------
CSS = """
<style>
.block-container{padding-top:1.4rem;max-width:1180px;}
.pg-hero{background:linear-gradient(120deg,#4f46e5 0%,#7c3aed 55%,#06b6d4 100%);border-radius:22px;
  padding:26px 32px;color:#fff;box-shadow:0 12px 32px rgba(79,70,229,.28);margin-bottom:14px;}
.pg-hero h1{margin:0;font-size:2.05rem;font-weight:800;color:#fff;letter-spacing:-.5px;padding:0;}
.pg-hero h1 .simo{display:inline-block;font-size:1.45em;font-weight:900;letter-spacing:1px;
  background:linear-gradient(90deg,#fff 0%,#a5f3fc 25%,#fde68a 50%,#a5f3fc 75%,#fff 100%);background-size:200% auto;
  -webkit-background-clip:text;background-clip:text;color:transparent;-webkit-text-fill-color:transparent;
  animation:pg-shine 3s linear infinite,pg-float 2.6s ease-in-out infinite;}
.pg-hero h1 .rest{display:inline-block;animation:pg-fade 1.2s ease both;}
@keyframes pg-shine{to{background-position:200% center;}}
@keyframes pg-float{0%,100%{transform:translateY(0);}50%{transform:translateY(-4px);}}
@keyframes pg-fade{from{opacity:0;transform:translateX(-10px);}to{opacity:1;transform:none;}}
.pg-hero p{margin:6px 0 0;opacity:.93;font-size:1.02rem;}
.pg-chip{display:inline-block;background:rgba(255,255,255,.18);padding:3px 12px;border-radius:999px;
  font-size:.78rem;margin:12px 6px 0 0;}
.pg-steps{display:flex;gap:10px;margin:0 0 10px;flex-wrap:wrap;}
.pg-step{flex:1;min-width:150px;border-radius:14px;padding:10px 14px;border:1px solid rgba(128,128,128,.25);
  background:rgba(128,128,128,.06);}
.pg-step b{display:block;font-size:.88rem;}
.pg-step span{font-size:.78rem;opacity:.75;}
.pg-step.done{border-color:rgba(16,185,129,.65);background:rgba(16,185,129,.11);}
.pg-step.active{border-color:rgba(79,70,229,.75);background:rgba(79,70,229,.11);box-shadow:0 0 0 3px rgba(79,70,229,.10);}
.pg-sec{display:flex;align-items:center;gap:12px;margin:24px 0 8px;}
.pg-sec .ic{width:40px;height:40px;border-radius:12px;display:flex;align-items:center;justify-content:center;
  font-size:1.2rem;background:linear-gradient(135deg,#4f46e5,#06b6d4);color:#fff;flex:none;}
.pg-sec h3{margin:0;padding:0;font-size:1.15rem;}
.pg-sec small{opacity:.7;}
.pg-badge{display:inline-block;padding:2px 11px;border-radius:999px;font-size:.74rem;font-weight:700;color:#fff;margin-right:6px;}
.pg-callout{border-left:4px solid #4f46e5;background:rgba(79,70,229,.08);padding:10px 14px;
  border-radius:0 12px 12px 0;margin:10px 0;line-height:1.5;}
.pg-callout.tip{border-left-color:#06b6d4;background:rgba(6,182,212,.09);}
.pg-empty{text-align:center;padding:38px 20px;border:2px dashed rgba(128,128,128,.35);border-radius:18px;margin:12px 0;}
.pg-empty .big{font-size:2.4rem;}
.pg-how{display:flex;gap:12px;flex-wrap:wrap;margin:6px 0 18px;}
.pg-how > div{flex:1;min-width:190px;padding:14px 16px;border-radius:16px;background:rgba(128,128,128,.07);
  border:1px solid rgba(128,128,128,.2);}
.pg-how .n{font-weight:800;color:#4f46e5;font-size:1.35rem;}
.pg-how b{display:block;margin-top:2px;}
.pg-how small{opacity:.75;}
.pg-side-title{font-weight:800;font-size:1.2rem;margin-bottom:6px;}
div[data-testid="stTabs"] button[role="tab"]{font-weight:600;padding:10px 18px;}
div[data-testid="stExpander"]{border-radius:14px;border:1px solid rgba(128,128,128,.25);box-shadow:0 2px 8px rgba(0,0,0,.04);}
div.stButton > button, div.stDownloadButton > button{border-radius:12px;font-weight:600;transition:all .15s ease;}
div.stButton > button:hover, div.stDownloadButton > button:hover{transform:translateY(-1px);box-shadow:0 6px 16px rgba(79,70,229,.25);}
div.stButton > button[kind="primary"]{background:linear-gradient(120deg,#4f46e5,#7c3aed);border:0;color:#fff;}
div[data-testid="stMetric"]{background:rgba(128,128,128,.07);border:1px solid rgba(128,128,128,.2);border-radius:14px;padding:12px 16px;}
div[data-testid="stFileUploader"] section{border-radius:16px;border:2px dashed rgba(79,70,229,.45);}
section[data-testid="stSidebar"]{border-right:1px solid rgba(128,128,128,.2);}
</style>
"""

CAT_STYLE = {  # kata kunci kategori -> (warna, ikon)
    "celah": ("#ef4444", "🕳️"), "data": ("#ef4444", "🕳️"), "metode": ("#3b82f6", "🧪"),
    "literatur": ("#8b5cf6", "📚"), "kebaruan": ("#f59e0b", "💡"), "konsisten": ("#14b8a6", "🔗"),
    "kontradiksi": ("#14b8a6", "🔗"), "validasi": ("#0ea5e9", "🎯"), "etika": ("#ec4899", "⚖️"),
}


def cat_style(kategori: str):
    k = kategori.lower()
    return next((v for key, v in CAT_STYLE.items() if key in k), ("#64748b", "📌"))


def esc(text: str) -> str:
    return html.escape(text or "").replace("\n", "<br>")


def sec(icon: str, title: str, desc: str = ""):
    st.markdown(
        f'<div class="pg-sec"><div class="ic">{icon}</div><div><h3>{html.escape(title)}</h3>'
        f"<small>{html.escape(desc)}</small></div></div>",
        unsafe_allow_html=True,
    )


def callout(text: str, tip: bool = False):
    st.markdown(f'<div class="pg-callout{" tip" if tip else ""}">{esc(text)}</div>', unsafe_allow_html=True)


def empty_state(icon: str, title: str, desc: str):
    st.markdown(
        f'<div class="pg-empty"><div class="big">{icon}</div><b>{html.escape(title)}</b><br>'
        f"<small>{html.escape(desc)}</small></div>",
        unsafe_allow_html=True,
    )


def wc(text: str) -> int:
    return len(re.findall(r"\w+", text or ""))


def steps_html() -> str:
    s = st.session_state
    n_notes = len(s.notes)
    n_done = sum(1 for n in s.notes if (s.fixes.get(n["id"]) or {}).get("hasil", "").strip())
    n_sec = sum(1 for v in s.sections.values() if v)
    total = len(s.section_names)
    steps = [
        ("📥 Bahan", f"{n_notes} catatan" if (s.inv or s.notes) else "belum dibaca", bool(s.inv or s.notes)),
        ("🧭 Rancang", f"{n_done}/{n_notes} ditangani" if n_notes else "menunggu draf", bool(n_notes) and n_done == n_notes),
        ("✍️ Tulis", f"{n_sec}/{total} bagian", n_sec > 0 and n_sec == total),
        ("✅ Review", "sudah direview" if s.qa.strip() else "belum", bool(s.qa.strip())),
    ]
    active = next((i for i, st_ in enumerate(steps) if not st_[2]), -1)
    cells = []
    for i, (title, sub, done) in enumerate(steps):
        cls = "done" if done else ("active" if i == active else "")
        cells.append(f'<div class="pg-step {cls}"><b>{title}{" ✓" if done else ""}</b><span>{html.escape(sub)}</span></div>')
    return f'<div class="pg-steps">{"".join(cells)}</div>'


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
st.set_page_config(page_title="SiMo SupeR PapeR", page_icon="📄", layout="wide")
init_state()
s = st.session_state
st.markdown(CSS, unsafe_allow_html=True)

with st.sidebar:
    st.markdown('<div class="pg-side-title">⚙️ Pengaturan</div>', unsafe_allow_html=True)
    with st.expander("🔑 API & model", expanded=True):
        s.api_key = st.text_input("Anthropic API key", value=s.api_key, type="password")
        s.model = st.text_input("Model", value=s.model)
    with st.expander("📝 Naskah", expanded=True):
        s.lang = st.selectbox("Bahasa naskah", ["Indonesia", "Inggris"], index=0 if s.lang == "Indonesia" else 1)
        s.ptype = st.selectbox("Jenis paper", list(PAPER_TYPES), index=list(PAPER_TYPES).index(s.ptype))
        s.target = st.text_input("Target jurnal/prosiding (opsional)", value=s.target)
        cites = ["APA 7", "Harvard", "IEEE", "Vancouver"]
        s.cite = st.selectbox("Gaya sitasi (pilih sebelum menulis)", cites, index=cites.index(s.cite))
    with st.expander("🔎 Sumber & batas"):
        use_web = st.checkbox("Izinkan pencarian web (literatur & sumber data)", value=True)
        per_doc = st.slider("Batas karakter per dokumen", 10_000, 200_000, 60_000, 10_000)
        total_cap = st.slider("Batas total konteks (karakter)", 30_000, 600_000, 200_000, 10_000)
    with st.expander("🧩 Skill"):
        skill_up = st.file_uploader("Ganti skill (.skill) — opsional", type=["skill", "zip"])
        s.skill, skill_info = load_skill(skill_up)
        if s.skill:
            st.caption(f"✅ Skill aktif: {skill_info}" + (" (unggahan)" if skill_up else " (bawaan)"))
        else:
            st.warning(f"Skill tidak aktif ({skill_info}). Aplikasi memakai prinsip dasar saja.")

s.section_names = PAPER_TYPES[s.ptype]

st.markdown(
    '<div class="pg-hero"><h1>📄 <span class="simo">SiMo</span> <span class="rest">SupeR PapeR</span></h1>'
    "<p>Ubah dokumen dan data mentah menjadi manuskrip jurnal — bertahap, transparan, dan bisa Anda edit di setiap langkah.</p>"
    f'<span class="pg-chip">🌐 {html.escape(s.lang)}</span>'
    f'<span class="pg-chip">📚 {html.escape(s.cite)}</span>'
    f'<span class="pg-chip">🧩 Skill {"aktif" if s.skill else "nonaktif"}</span>'
    f'<span class="pg-chip">🤖 {html.escape(s.model)}</span></div>',
    unsafe_allow_html=True,
)
steps_ph = st.empty()  # diisi di akhir skrip agar mencerminkan state terbaru

tab1, tab2, tab3, tab4, tab5 = st.tabs(
    ["📥 1 · Bahan", "🧭 2 · Rancang", "✍️ 3 · Tulis", "✅ 4 · Review & Ekspor", "🎤 5 · Buat presentasi"]
)

# --- Tab 1: Bahan -----------------------------------------------------------
with tab1:
    st.markdown(
        '<div class="pg-how">'
        '<div><span class="n">1</span><b>Unggah</b><small>Satu atau beberapa dokumen/data.</small></div>'
        '<div><span class="n">2</span><b>Baca dokumen</b><small>Skill menyusun inventaris, klaim inti, kerangka, dan catatan.</small></div>'
        '<div><span class="n">3</span><b>Lanjut ke Rancang</b><small>Tangani catatan dengan dokumen/data tambahan.</small></div>'
        "</div>",
        unsafe_allow_html=True,
    )
    files = st.file_uploader(
        "Unggah satu atau beberapa dokumen/data (PDF, DOCX, XLSX, CSV, TXT, MD)",
        type=FILE_TYPES,
        accept_multiple_files=True,
    )
    if st.button("🔍 Baca dokumen", type="primary", disabled=not files):
        if not s.api_key:
            st.error("Isi API key di sidebar terlebih dahulu.")
        elif analyze_documents(files, per_doc, total_cap):
            st.success(f"Draf tersimpan: {len(s.notes)} catatan perbaikan. Lanjutkan di tab 2 (Rancang).")
    elif s.inv or s.notes:
        st.caption(f"✅ Draf bahan tersimpan · {len(s.notes)} catatan · lanjutkan di tab 2 (Rancang).")
    if s.inv or s.notes:
        m1, m2, m3 = st.columns(3)
        m1.metric("Berkas terbaca", s.ctx.count("=== [DOK-"))
        m2.metric("Karakter bahan", f"{len(s.ctx):,}")
        m3.metric("Catatan perbaikan", len(s.notes))

# --- Tab 2: Rancang ---------------------------------------------------------
with tab2:
    if not (s.inv or s.claim or s.notes):
        empty_state("🧭", "Belum ada draf", "Unggah dokumen lalu klik “Baca dokumen” di tab 1.")
    else:
        sec("🗂️", "Draf dari tab Bahan", "Inventaris, klaim inti, dan kerangka — bisa diedit.")
        with st.expander("Buka / edit draf", expanded=False):
            editor("Inventaris", "inv", 220)
            editor("Klaim inti", "claim", 180)
            editor("Kerangka", "outline", 220)

        n_open = sum(1 for n in s.notes if not (s.fixes.get(n["id"]) or {}).get("hasil", "").strip())
        sec("🛠️", "Catatan perbaikan", "Untuk tiap catatan: unggah dokumen lain, atau cari data di sumber terpercaya.")
        k1, k2, k3 = st.columns(3)
        k1.metric("Total catatan", len(s.notes))
        k2.metric("Terbuka", n_open)
        k3.metric("Sudah ditangani", len(s.notes) - n_open)
        for n in s.notes:
            nid = n["id"]
            fx = s.fixes.get(nid)
            handled = bool(fx and fx.get("hasil", "").strip())
            color, icon = cat_style(n["kategori"])
            with st.expander(f"{'✅' if handled else '⚠️'} #{nid} {icon} {n['judul']}", expanded=not handled):
                status = (
                    '<span class="pg-badge" style="background:#10b981">Ditangani</span>' if handled
                    else '<span class="pg-badge" style="background:#f59e0b">Terbuka</span>'
                )
                st.markdown(
                    f'<span class="pg-badge" style="background:{color}">{esc(n["kategori"])}</span>{status}',
                    unsafe_allow_html=True,
                )
                callout(n["detail"])
                callout("Saran: " + n["saran"], tip=True)
                c1, c2 = st.columns(2)
                with c1:
                    ups = st.file_uploader(
                        "Opsi 1 — Upload dokumen lain", type=FILE_TYPES, accept_multiple_files=True,
                        key=f"up_{nid}_{s.ver.get(f'up_{nid}', 0)}",
                    )
                    if st.button("📎 Gunakan dokumen ini", key=f"useup_{nid}", disabled=not ups):
                        if need_ready():
                            handle_upload_fix(n, ups, per_doc)
                            st.rerun()
                with c2:
                    st.markdown("**Opsi 2 — Cari data di sumber terpercaya**")
                    st.caption("Jurnal bereputasi, lembaga resmi, organisasi internasional.")
                    if st.button("🌐 Cari di sumber terpercaya", key=f"web_{nid}", disabled=not use_web):
                        if need_ready():
                            handle_web_fix(n)
                            st.rerun()
                if fx:
                    st.caption(f"Ditangani lewat: {fx['sumber']}")
                    fx["hasil"] = st.text_area(
                        "Hasil (bisa diedit; dipakai saat menulis)", value=fx["hasil"], height=200,
                        key=f"fx_{nid}_{s.ver.get(f'fx_{nid}', 0)}",
                    )

# --- Tab 3: Tulis -----------------------------------------------------------
with tab3:
    sec("✍️", "Susun paper lengkap", "Dari draf tab 1 dan data tambahan di tab 2.")
    written = [n for n in s.section_names if s.sections.get(n)]
    w1, w2, w3 = st.columns(3)
    w1.metric("Bagian ditulis", f"{len(written)}/{len(s.section_names)}")
    w2.metric("Total kata", f"{sum(wc(s.sections[n]) for n in written):,}")
    w3.metric("Gambar", len(s.figures))
    if st.button("🚀 Tulis semua bagian", type="primary"):
        if need_ready():
            if use_web and not s.lit:
                search_literature()
            bar = st.progress(0.0)
            for idx, name in enumerate(s.section_names):
                with st.spinner(f"Menulis {name}..."):
                    write_section(name)
                bar.progress((idx + 1) / len(s.section_names))
            st.rerun()

    for name in s.section_names:
        with st.expander(f"{'✅' if s.sections.get(name) else '⬜'} {name}", expanded=bool(s.sections.get(name))):
            key = f"sec_{name}"
            s.sections[name] = st.text_area(
                name, value=s.sections.get(name, ""), height=320, key=f"ta_{key}_{s.ver.get(key, 0)}",
                label_visibility="collapsed",
            )
            if st.button(f"🔄 Tulis ulang '{name}'", key=f"re_{name}"):
                if need_ready():
                    with st.spinner(f"Menulis ulang {name}..."):
                        write_section(name)
                    st.rerun()

    with st.expander("📚 Catatan literatur (dicari otomatis saat 'Tulis semua bagian'; bisa diedit/ditempel)"):
        editor("Literatur", "lit", 240)
        if st.button("Cari ulang literatur", disabled=not use_web):
            if need_ready():
                search_literature()
                st.rerun()

    sec("📊", "Grafik dari data", "Dibuat dari tabel CSV/XLSX pada dokumen Anda — tanpa angka rekaan.")
    if not s.tables:
        empty_state("📊", "Belum ada tabel data", "Unggah CSV/XLSX di dokumen utama atau dokumen tambahan.")
    else:
        g1, g2 = st.columns(2)
        with g1:
            tname = st.selectbox("Tabel data", list(s.tables))
            df = s.tables[tname]
            kind = st.selectbox("Jenis grafik", ["Garis", "Batang", "Sebaran"])
        with g2:
            x = st.selectbox("Sumbu X", list(df.columns))
            numcols = [c for c in df.select_dtypes("number").columns if c != x]
            ys = st.multiselect("Sumbu Y (kolom numerik)", numcols)
        cap = st.text_input("Keterangan gambar", value=(f"{', '.join(map(str, ys))} menurut {x}" if ys else ""))
        if st.button("➕ Tambah grafik", disabled=not ys):
            s.figures.append({"png": make_chart(df, kind, x, ys), "caption": cap})
            st.rerun()
    for i, fig in enumerate(s.figures, 1):
        st.image(fig["png"], caption=f"Gambar {i}. {fig['caption']}", width=480)
        if st.button(f"🗑️ Hapus Gambar {i}", key=f"delfig_{i}"):
            s.figures.pop(i - 1)
            st.rerun()

    sec("🏷️", "Judul, abstrak, kata kunci", "Dibuat setelah semua bagian ditulis.")
    if st.button("Buat judul & abstrak (setelah semua bagian ditulis)"):
        if need_ready(need_text=True):
            make_meta()
            st.rerun()
    if s.meta:
        s.meta["judul"] = st.text_input("Judul", value=s.meta.get("judul", ""))
        s.meta["abstrak"] = st.text_area("Abstrak", value=s.meta.get("abstrak", ""), height=200)
        s.meta["kata_kunci"] = st.text_input("Kata kunci", value=s.meta.get("kata_kunci", ""))

# --- Tab 4: Review & Ekspor -------------------------------------------------
with tab4:
    sec("🔎", "1 · Review hasil tulis", "Dinilai seperti reviewer jurnal bereputasi.")
    if st.button("Review naskah"):
        if need_ready(need_text=True):
            run_review()
            st.rerun()
    editor("Hasil review (bisa diedit sebelum dipakai untuk perbaikan)", "qa", 320)

    sec("🩹", "2 · Perbaiki paper sesuai review", "Tiap bagian diperbaiki; versi lama bisa dikembalikan.")
    c1, c2 = st.columns([2, 1])
    if c1.button("✨ Perbaiki paper sesuai hasil review", type="primary", disabled=not s.qa.strip()):
        if need_ready(need_text=True):
            revise_all()
            st.rerun()
    if s.backup and c2.button("↩️ Kembalikan versi sebelum perbaikan"):
        s.sections, s.meta, s.refs = dict(s.backup["sections"]), dict(s.backup["meta"]), s.backup["refs"]
        s.backup = None
        for n in s.section_names:
            bump(f"sec_{n}")
        bump("refs")
        st.rerun()
    st.caption("Setelah perbaikan, jalankan review lagi bila perlu, lalu susun ulang daftar pustaka.")

    sec("📖", "3 · Daftar pustaka", "Hanya dari sitasi yang ada di naskah.")
    if st.button("Susun daftar pustaka dari sitasi di naskah"):
        if need_ready(need_text=True):
            body = "\n\n".join(f"## {n}\n{s.sections.get(n, '')}" for n in s.section_names)
            with st.spinner("Menyusun daftar pustaka..."):
                refs = call_llm(
                    f"NASKAH:\n{body}\n\nCATATAN LITERATUR:\n{s.lit}\n\n{additions_block()}\n\n"
                    f"Susun daftar pustaka HANYA untuk sumber yang disitasi di naskah dengan gaya {s.cite} "
                    "(ikuti urutan gaya itu: abjad untuk APA/Harvard, urutan kemunculan untuk IEEE/Vancouver). "
                    "Gunakan data dari catatan literatur dan data tambahan; jangan menambah detail yang tidak ada. "
                    "Pertahankan penanda [VERIFIKASI]. Di akhir, daftar sitasi di naskah yang tidak ada di sumber "
                    "dengan judul 'Sitasi tanpa sumber'.",
                    max_tokens=4000,
                )
            if refs:
                s.refs = refs
                bump("refs")
                st.rerun()
    editor("Daftar pustaka (bisa diedit)", "refs", 220)

    sec("📦", "4 · Ekspor", "Unduh naskah final.")
    manuscript = assemble_manuscript()
    d1, d2 = st.columns(2)
    d1.download_button(
        "⬇️ Unduh .docx", data=md_to_docx(manuscript, s.figures), file_name="manuskrip.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )
    d2.download_button(
        "⬇️ Unduh .md", data=figs_to_text(manuscript), file_name="manuskrip.md", mime="text/markdown",
        use_container_width=True,
    )
    st.caption("Periksa semua penanda [DATA DIBUTUHKAN] dan [VERIFIKASI] sebelum submit ke jurnal.")

# ----------------------------------------------------------------------------
# Tab 5 · Buat presentasi (skill presentasi-dari-dokumen)
# ----------------------------------------------------------------------------
PRES_SKILL_DIR = os.path.join(APP_DIR, "skill_presentasi")

PRES_BASE = """Anda adalah perancang presentasi ilmiah/profesional yang membuat bahan presentasi dari naskah paper final pengguna.

PRINSIP WAJIB:
1. Isi hanya dari sumber (naskah final dan dokumen rujukan). Jangan mengarang angka, nama, tanggal, atau klaim. Bila data tidak ada, tulis "[data belum tersedia]".
2. Rujukan mengikat bentuk, sumber mengisi isi. Bila bertentangan, ikuti rujukan dan laporkan konfliknya.
3. Satu slide, satu pesan; judul slide berupa kalimat kesimpulan, bukan label.
4. Visual (angka besar, grafik, tabel, gambar) menggantikan paragraf; narasi panjang dipindah ke catatan pembicara.
5. Bahasa mengikuti bahasa naskah."""

PRES_APP_NOTE = (
    "\n\nCATATAN PENERAPAN: Anda dipanggil oleh aplikasi, bukan percakapan. Jangan bertanya balik dan jangan menulis "
    "basa-basi. File .pptx dibuat oleh kode aplikasi (python-pptx), bukan oleh Anda: abaikan bagian skill yang "
    "menyebut skill pptx/pptxgenjs, perintah bash, jalur /mnt, render gambar, markitdown, atau present_files. "
    "Kerjakan tugas pada pesan pengguna (brief, fakta, rencana slide, QA) dan patuhi format keluaran yang diminta."
)


def load_pres_skill():
    """Kembalikan (teks_skill, keterangan) dari folder ./skill_presentasi."""
    docs = []
    try:
        for root, _, fs in os.walk(PRES_SKILL_DIR):
            for fn in sorted(fs):
                if fn.endswith(".md"):
                    p = os.path.join(root, fn)
                    with open(p, encoding="utf-8") as fh:
                        docs.append((os.path.relpath(p, PRES_SKILL_DIR), fh.read()))
    except Exception as e:  # noqa: BLE001
        return "", f"gagal memuat skill: {e}"
    main = [d for d in docs if d[0].endswith("SKILL.md")]
    if not main:
        return "", "SKILL.md tidak ditemukan di skill_presentasi/"
    text = _strip_frontmatter(main[0][1])
    for p, t in docs:
        if not p.endswith("SKILL.md"):
            text += f"\n\n--- Referensi skill: {p} ---\n{t}"
    return text, f"{len(text):,} karakter"


def call_llm_pres(user: str, max_tokens: int = 6000) -> str:
    s = st.session_state
    system = [{
        "type": "text",
        "text": PRES_BASE + f"\n\n=== SKILL: presentasi-dari-dokumen (ikuti alur dan prinsipnya) ===\n{s.pres_skill}" + PRES_APP_NOTE,
        "cache_control": {"type": "ephemeral"},
    }]
    try:
        resp = Anthropic(api_key=s.api_key).messages.create(
            model=s.model, max_tokens=max_tokens, system=system, messages=[{"role": "user", "content": user}]
        )
    except Exception as e:  # noqa: BLE001
        st.error(f"Gagal memanggil model: {e}")
        return ""
    return "".join(b.text for b in resp.content if getattr(b, "type", "") == "text").strip()


def pptx_summary(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = ["Layout tersedia: " + ", ".join(lay.name for lay in prs.slide_layouts)]
    for i, sl in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text.strip() for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        out.append(f"Slide {i}: " + " | ".join(texts)[:400])
    return "\n".join(out)


def parse_plan(text: str) -> dict:
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", (text or "").strip(), flags=re.M).strip()
    a, b = t.find("{"), t.rfind("}")
    data = json.loads(t[a: b + 1])
    if not isinstance(data.get("slides"), list) or not data["slides"]:
        raise ValueError("rencana tidak memuat daftar 'slides'")
    return data


def _fmt_num(v) -> str:
    try:
        return format(float(v), "f").rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        return str(v)


_NUM = re.compile(r"\d[\d.,]*\d|\d")


def check_plan(plan: dict, source: str) -> list:
    """Pemeriksaan deterministik: angka tanpa dasar di sumber, slide terlalu padat, pembuka/penutup."""
    toks, digs = set(), set()
    for m in _NUM.findall(source or ""):
        t = m.rstrip(".,")
        toks.add(t)
        digs.add(re.sub(r"\D", "", t))
    issues, slides = [], plan.get("slides", [])
    for i, sl in enumerate(slides, 1):
        parts = [str(sl.get("judul", ""))] + [str(x) for x in (sl.get("poin") or [])]
        a = sl.get("angka") or {}
        parts += [str(a.get("nilai", "")), str(a.get("label", ""))]
        for it in (sl.get("bagan") or {}).get("seri") or []:
            parts += [str(it.get("label", "")), _fmt_num(it.get("nilai"))]
        tb = sl.get("tabel") or {}
        parts += [str(x) for x in (tb.get("header") or [])]
        for r in tb.get("baris") or []:
            parts += [str(x) for x in r]
        seen = set()
        for m in _NUM.findall(" ".join(parts)):
            t = m.rstrip(".,")
            d = re.sub(r"\D", "", t)
            if len(d) <= 1 or t in seen:
                continue
            seen.add(t)
            if t not in toks and d not in digs:
                issues.append(f"- [Akurasi] Slide {i}: angka '{t}' tidak ditemukan di fakta kunci/naskah")
        n_words = sum(len(re.findall(r"\w+", str(x))) for x in (sl.get("poin") or []))
        if sl.get("tipe") in ("isi", "angka", "bagan", "tabel", "gambar") and n_words > 40:
            issues.append(f"- [Ringkas] Slide {i}: {n_words} kata pada poin (batas ±40)")
    if slides and slides[0].get("tipe") != "judul":
        issues.append("- [Struktur] Slide pertama bukan slide judul/pembuka")
    if slides and slides[-1].get("tipe") != "penutup":
        issues.append("- [Struktur] Slide terakhir bukan slide penutup/kesimpulan")
    return issues


def build_pptx(plan: dict, figures: list, template=None):
    """Bangun .pptx dari rencana slide. Kembalikan (bytes, daftar_peringatan)."""
    from PIL import Image
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    warns = []
    use_tpl = template is not None
    if use_tpl:
        prs = Presentation(io.BytesIO(template))
        for sld in list(prs.slides._sldIdLst):  # kosongkan slide contoh, pertahankan master/layout
            prs.part.drop_rel(sld.rId)
            prs.slides._sldIdLst.remove(sld)
        W, H = prs.slide_width / 914400, prs.slide_height / 914400
        P = {"gelap": MSO_THEME_COLOR.TEXT_1, "utama": MSO_THEME_COLOR.ACCENT_1, "aksen": MSO_THEME_COLOR.ACCENT_2,
             "terang": MSO_THEME_COLOR.BACKGROUND_1, "abu": MSO_THEME_COLOR.TEXT_2, "putih": MSO_THEME_COLOR.BACKGROUND_1}
        fj = fi = None
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        W, H = 13.333, 7.5
        pal = plan.get("palet") or {}

        def _hex(v, default):
            m = re.fullmatch(r"#?([0-9A-Fa-f]{6})", str(v or "").strip())
            return RGBColor.from_string((m.group(1) if m else default).upper())

        P = {"gelap": _hex(pal.get("gelap"), "1F2937"), "utama": _hex(pal.get("utama"), "0F766E"),
             "aksen": _hex(pal.get("aksen"), "F59E0B"), "terang": _hex(pal.get("terang"), "F8FAFC"),
             "abu": RGBColor(0x6B, 0x72, 0x80), "putih": RGBColor(0xFF, 0xFF, 0xFF)}
        ok_fonts = {"Calibri", "Arial", "Georgia", "Cambria", "Segoe UI", "Trebuchet MS", "Verdana", "Tahoma", "Garamond"}
        fj = plan.get("font_judul") if plan.get("font_judul") in ok_fonts else "Georgia"
        fi = plan.get("font_isi") if plan.get("font_isi") in ok_fonts else "Calibri"

    def setc(cf, c):
        if isinstance(c, RGBColor):
            cf.rgb = c
        else:
            cf.theme_color = c

    def add_text(sl, x, y, w, h, text, size, color, bold=False, font=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, line in enumerate(text if isinstance(text, list) else [text]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = str(line)
            r.font.size, r.font.bold = Pt(size), bold
            if font:
                r.font.name = font
            setc(r.font.color, color)
        return tb

    def add_bullets(sl, x, y, w, h, items, size=20):
        tb = add_text(sl, x, y, w, h, [str(i) for i in items], size, P["gelap"], font=fi)
        for p in tb.text_frame.paragraphs:
            p.space_after = Pt(10)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "342900")
            pPr.set("indent", "-285750")
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        return tb

    def add_rect(sl, x, y, w, h, color):
        shp = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        setc(shp.fill.fore_color, color)
        shp.line.fill.background()
        return shp

    def bg(sl, color):
        if not use_tpl:
            sl.background.fill.solid()
            setc(sl.background.fill.fore_color, color)

    title_types = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    skip_types = title_types + (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)

    def pick_content_layout():
        best, best_n = None, 99
        for lay in prs.slide_layouts:
            if not any(p.placeholder_format.type in title_types for p in lay.placeholders):
                continue
            n_other = sum(1 for p in lay.placeholders if p.placeholder_format.type not in skip_types)
            if "title only" in lay.name.lower() or "judul saja" in lay.name.lower():
                return lay
            if n_other < best_n:
                best, best_n = lay, n_other
        return best or prs.slide_layouts[0]

    def clean_placeholders(sl, keep=()):
        for ph in list(sl.placeholders):
            if ph.placeholder_format.idx in keep or ph.placeholder_format.type in title_types:
                continue
            if not (ph.has_text_frame and ph.text_frame.text.strip()):
                ph._element.getparent().remove(ph._element)

    figures = figures or []
    slides = plan.get("slides", [])
    total = len(slides)
    for n, sd in enumerate(slides, 1):
        tipe = str(sd.get("tipe") or "isi").lower()
        judul = str(sd.get("judul") or "")
        poin = [str(x) for x in (sd.get("poin") or [])]
        sumber = str(sd.get("sumber") or "")
        # validasi data khusus; bila tidak lengkap, jatuh ke slide teks biasa
        bagan = sd.get("bagan") or {}
        seri = [it for it in (bagan.get("seri") or []) if isinstance(it, dict) and str(it.get("label", "")).strip()]
        try:
            seri = [{"label": str(it["label"]), "nilai": float(it["nilai"])} for it in seri]
        except (TypeError, ValueError, KeyError):
            seri = []
        tabel = sd.get("tabel") or {}
        header = [str(x) for x in (tabel.get("header") or [])]
        baris = [[str(c) for c in r] for r in (tabel.get("baris") or []) if isinstance(r, (list, tuple))][:8]
        gi = sd.get("gambar")
        gi = gi if isinstance(gi, int) and 1 <= gi <= len(figures) else None
        angka = sd.get("angka") or {}
        if tipe == "bagan" and len(seri) < 2:
            warns.append(f"Slide {n}: data bagan kurang dari 2 titik; dibuat sebagai slide teks.")
            tipe = "isi"
        if tipe == "tabel" and not (header and baris):
            warns.append(f"Slide {n}: data tabel tidak lengkap; dibuat sebagai slide teks.")
            tipe = "isi"
        if tipe == "gambar" and gi is None:
            warns.append(f"Slide {n}: nomor gambar tidak valid; dibuat sebagai slide teks.")
            tipe = "isi"
        if tipe == "angka" and not str(angka.get("nilai", "")).strip():
            warns.append(f"Slide {n}: nilai angka kosong; dibuat sebagai slide teks.")
            tipe = "isi"

        if tipe in ("judul", "penutup"):
            if use_tpl:
                sl = prs.slides.add_slide(prs.slide_layouts[0])
                if sl.shapes.title is not None:
                    sl.shapes.title.text = judul
                sub = next((p for p in sl.placeholders if p.placeholder_format.idx == 1), None)
                if sub is not None and (poin or sd.get("subjudul")):
                    sub.text = "\n".join(poin) if poin else str(sd.get("subjudul"))
                clean_placeholders(sl, keep=(1,))
            else:
                sl = prs.slides.add_slide(prs.slide_layouts[6])
                bg(sl, P["gelap"])
                add_rect(sl, 0.8, 2.55, 1.2, 0.08, P["aksen"])
                add_text(sl, 0.8, 2.75, W - 1.6, 1.8, judul, 40, P["putih"], True, fj, anchor=MSO_ANCHOR.TOP)
                sub = "\n".join(poin) if poin else str(sd.get("subjudul") or plan.get("subjudul") or "")
                if sub:
                    add_text(sl, 0.8, 4.6, W - 1.6, 1.8, sub.split("\n"), 20, P["terang"], False, fi)
        else:
            if use_tpl:
                sl = prs.slides.add_slide(pick_content_layout())
                if sl.shapes.title is not None:
                    sl.shapes.title.text = judul
                    t = sl.shapes.title
                    top = (t.top + t.height) / 914400 + 0.15
                else:
                    top = 1.5
                clean_placeholders(sl)
                left = 0.6
            else:
                sl = prs.slides.add_slide(prs.slide_layouts[6])
                bg(sl, P["terang"])
                add_rect(sl, 0, 0.45, 0.14, 1.0, P["utama"])
                add_text(sl, 0.55, 0.35, W - 1.1, 1.2, judul, 28, P["gelap"], True, fj, anchor=MSO_ANCHOR.MIDDLE)
                top, left = 1.75, 0.6
            bw, bh = W - 2 * left, H - top - 0.75

            if tipe == "angka":
                add_text(sl, left, top + 0.3, bw * 0.5, 2.0, str(angka.get("nilai")), 88, P["utama"], True, fj)
                add_text(sl, left, top + 2.4, bw * 0.5, 1.6, str(angka.get("label", "")), 24, P["gelap"], False, fi)
                if poin:
                    add_bullets(sl, left + bw * 0.55, top + 0.3, bw * 0.45, bh - 0.3, poin, 20)
            elif tipe == "bagan":
                cd = CategoryChartData()
                cd.categories = [it["label"] for it in seri]
                cd.add_series(str(bagan.get("satuan") or "Nilai"), [it["nilai"] for it in seri])
                jenis = str(bagan.get("jenis") or "batang").lower()
                ctype = {"garis": XL_CHART_TYPE.LINE_MARKERS, "pai": XL_CHART_TYPE.PIE}.get(jenis, XL_CHART_TYPE.COLUMN_CLUSTERED)
                cw = bw * (0.62 if poin else 1.0)
                ch = sl.shapes.add_chart(ctype, Inches(left), Inches(top), Inches(cw), Inches(bh), cd).chart
                ch.font.size = Pt(14)
                ch.has_legend = jenis == "pai"
                if ch.has_legend:
                    ch.legend.position = XL_LEGEND_POSITION.RIGHT
                    ch.legend.include_in_layout = False
                ch.plots[0].has_data_labels = True
                ch.plots[0].data_labels.font.size = Pt(14)
                if not use_tpl and jenis != "pai":
                    ser = ch.plots[0].series[0]
                    if jenis == "garis":
                        ser.format.line.color.rgb = P["utama"]
                    else:
                        ser.format.fill.solid()
                        ser.format.fill.fore_color.rgb = P["utama"]
                if poin:
                    add_bullets(sl, left + cw + 0.3, top, bw - cw - 0.3, bh, poin, 20)
            elif tipe == "tabel":
                rows = [header] + baris
                ncol = len(header)
                tb = sl.shapes.add_table(len(rows), ncol, Inches(left), Inches(top), Inches(bw), Inches(0.5 * len(rows))).table
                for r, row in enumerate(rows):
                    for c in range(ncol):
                        cell = tb.cell(r, c)
                        cell.text = row[c] if c < len(row) else ""
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size, run.font.bold = Pt(16), r == 0
                                if fi:
                                    run.font.name = fi
                                if r == 0 and not use_tpl:
                                    run.font.color.rgb = P["putih"]
                        if r == 0 and not use_tpl:
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = P["utama"]
                if poin:
                    add_bullets(sl, left, top + 0.5 * len(rows) + 0.3, bw, max(0.8, bh - 0.5 * len(rows) - 0.3), poin, 18)
            elif tipe == "gambar":
                png = figures[gi - 1]["png"]
                iw, ih = Image.open(io.BytesIO(png)).size
                box_w = bw * (0.62 if poin else 1.0)
                box_h = bh - 0.45
                sc = min(box_w / iw, box_h / ih)
                pw, ph_ = iw * sc, ih * sc
                sl.shapes.add_picture(io.BytesIO(png), Inches(left + (box_w - pw) / 2), Inches(top), Inches(pw), Inches(ph_))
                add_text(sl, left, top + ph_ + 0.05, box_w, 0.4, f"Gambar {gi}. {figures[gi - 1]['caption']}", 12, P["abu"],
                         False, fi, PP_ALIGN.CENTER)
                if poin:
                    add_bullets(sl, left + box_w + 0.3, top, bw - box_w - 0.3, bh, poin, 20)
            else:  # isi
                if poin:
                    add_bullets(sl, left, top, bw, bh, poin, 22)
            if sumber:
                add_text(sl, left, H - 0.6, bw - 1.0, 0.4, "Sumber: " + sumber, 11, P["abu"], False, fi)
            add_text(sl, W - 1.3, H - 0.6, 0.7, 0.4, str(n), 11, P["abu"], False, fi, PP_ALIGN.RIGHT)
        notes = str(sd.get("catatan") or "").strip()
        if notes:
            sl.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), warns


PLAN_SCHEMA = """{
 "judul_deck": "...", "subjudul": "...",
 "palet": {"gelap": "1F2937", "utama": "0F766E", "aksen": "F59E0B", "terang": "F8FAFC"},
 "font_judul": "Georgia", "font_isi": "Calibri",
 "slides": [
  {"tipe": "judul|isi|angka|bagan|tabel|gambar|penutup", "judul": "kalimat kesimpulan",
   "poin": ["maks. 3-5 poin singkat"],
   "angka": {"nilai": "32%", "label": "satu baris konteks"},
   "bagan": {"jenis": "batang|garis|pai", "satuan": "satuan", "seri": [{"label": "...", "nilai": 0}]},
   "tabel": {"header": ["..."], "baris": [["..."]]},
   "gambar": 1,
   "sumber": "Naskah, bagian ...", "catatan": "narasi 2-4 kalimat + rujukan fakta"}
 ]
}"""


def generate_presentation(manuscript: str, ref_files, audiens: str, tujuan: str, durasi: int) -> bool:
    s = st.session_state
    s.pres_skill, info = load_pres_skill()
    if not s.pres_skill:
        st.error(f"Skill presentasi tidak termuat ({info}).")
        return False
    template, chunks = None, []
    for f in ref_files or []:
        data = f.getvalue()
        if f.name.lower().endswith(".pptx"):
            template = template or data
            try:
                chunks.append(f"=== [RUJUKAN] {f.name} (pptx) ===\n{pptx_summary(data)}")
            except Exception as e:  # noqa: BLE001
                chunks.append(f"=== [RUJUKAN] {f.name} (pptx) ===\n[GAGAL MEMBACA: {e}]")
                template = None if template is data else template
        else:
            chunks.append(f"=== [RUJUKAN] {f.name} ===\n{extract_text(f)[:60_000]}")
    ref_text = "\n\n".join(chunks)[:150_000]

    brief = "(tidak ada dokumen rujukan; gunakan pola default skill untuk laporan hasil kajian)"
    if ref_text:
        with st.spinner("Langkah 2 · Mengekstrak brief rujukan..."):
            brief = call_llm_pres(
                f"DOKUMEN RUJUKAN:\n{ref_text}\n\nTUGAS: jalankan Langkah 2 skill — tulis brief rujukan memakai format "
                "pada references/brief-rujukan.md (8 butir; tulis 'tidak diatur' bila rujukan diam; sertakan lokasi "
                "aturan).", max_tokens=3000)
        if not brief:
            return False
    with st.spinner("Langkah 3 · Mengekstrak fakta kunci dari naskah..."):
        fakta = call_llm_pres(
            f"NASKAH FINAL:\n{manuscript}\n\nTUGAS: jalankan Langkah 3 skill — daftar fakta kunci (temuan, angka, tabel "
            "sebagai baris, metode ringkas, kontribusi, keterbatasan, kesimpulan). Setiap butir diberi penanda sumber "
            "(Naskah, bagian ...). Ini satu-satunya basis isi slide.", max_tokens=5000)
    if not fakta:
        return False
    figs = "\n".join(f"Gambar {i}: {f['caption']}" for i, f in enumerate(s.figures, 1)) or "(tidak ada)"
    plan_prompt = (
        f"AUDIENS: {audiens}\nTUJUAN: {tujuan}\nDURASI: {durasi} menit (sekitar 1 slide per 1,5-2 menit; ikuti batas "
        f"jumlah slide pada brief rujukan bila ada)\nBahasa deck: {s.lang}\n\nBRIEF RUJUKAN:\n{brief}\n\n"
        f"FAKTA KUNCI:\n{fakta}\n\nGAMBAR TERSEDIA dari naskah (untuk slide tipe 'gambar', isi field gambar dengan nomornya):\n"
        f"{figs}\n\nTUGAS: jalankan Langkah 4-6 skill — rancang storyline dan isi tiap slide. Keluarkan HANYA JSON valid "
        f"(tanpa ``` dan tanpa teks lain) dengan skema:\n{PLAN_SCHEMA}\n\nATURAN: judul slide = kalimat kesimpulan; maks. "
        "±40 kata poin per slide isi; angka hanya dari FAKTA KUNCI; bagan hanya dari angka di fakta; slide pembuka "
        "('judul') dan penutup ('penutup') wajib kecuali rujukan melarang; isi field hanya yang relevan per tipe; "
        "'sumber' memuat bagian naskah/dokumen rujukan; 'catatan' berisi narasi bicara 2-4 kalimat + rujukan fakta; "
        "bila data kurang tulis \"[data belum tersedia]\"."
    )
    with st.spinner("Langkah 4-6 · Merancang storyline dan isi slide..."):
        raw = call_llm_pres(plan_prompt, max_tokens=8000)
    if not raw:
        return False
    try:
        plan = parse_plan(raw)
    except Exception:  # noqa: BLE001
        with st.spinner("Merapikan format rencana slide..."):
            fixed = call_llm_pres(f"Perbaiki teks berikut menjadi JSON valid sesuai skema, keluarkan HANYA JSON:\n{raw}", 8000)
        try:
            plan = parse_plan(fixed)
        except Exception as e:  # noqa: BLE001
            st.error(f"Rencana slide tidak bisa dibaca sebagai JSON: {e}")
            return False

    source = fakta + "\n" + manuscript
    with st.spinner("Langkah 7 · QA kepatuhan rujukan dan akurasi..."):
        det = check_plan(plan, source)
        llm_qa = call_llm_pres(
            f"BRIEF RUJUKAN:\n{brief}\n\nFAKTA KUNCI:\n{fakta}\n\nRENCANA SLIDE:\n{json.dumps(plan, ensure_ascii=False)}\n\n"
            "TUGAS: jalankan QA Langkah 7 — (1) kepatuhan terhadap brief rujukan butir demi butir, (2) akurasi: klaim/"
            "angka tanpa dasar di fakta, (3) judul slide yang masih berupa label, bukan kesimpulan. Tulis satu baris per "
            "masalah: '- [Kategori] Slide n: masalah'. Bila tidak ada masalah tulis persis: TIDAK ADA MASALAH.",
            max_tokens=2500)
    has_llm_issue = bool(llm_qa) and "TIDAK ADA MASALAH" not in llm_qa.upper()
    if det or has_llm_issue:
        with st.spinner("Memperbaiki rencana slide sesuai QA..."):
            fixed = call_llm_pres(
                f"BRIEF RUJUKAN:\n{brief}\n\nFAKTA KUNCI:\n{fakta}\n\nRENCANA SLIDE SAAT INI:\n"
                f"{json.dumps(plan, ensure_ascii=False)}\n\nMASALAH QA:\n" + "\n".join(det) + "\n" + (llm_qa if has_llm_issue else "")
                + "\n\nTUGAS: perbaiki rencana hanya pada bagian yang bermasalah (angka tanpa dasar dihapus atau diganti "
                "'[data belum tersedia]', poin diringkas, judul dijadikan kesimpulan). Keluarkan rencana lengkap sebagai "
                f"HANYA JSON valid dengan skema yang sama:\n{PLAN_SCHEMA}", max_tokens=8000)
        try:
            plan = parse_plan(fixed)
        except Exception:  # noqa: BLE001
            st.warning("Perbaikan otomatis gagal dibaca; memakai rencana awal.")
    remaining = check_plan(plan, source)

    try:
        pptx_bytes, warns = build_pptx(plan, s.figures, template)
    except Exception as e:  # noqa: BLE001
        st.error(f"Gagal membangun .pptx: {e}")
        return False
    qa = ["### Temuan awal", *(det or ["- Pemeriksaan otomatis: tidak ada masalah."]),
          llm_qa or "", "", "### Setelah perbaikan otomatis",
          *(remaining or ["- Tidak ada angka tanpa dasar, slide padat, atau masalah struktur."]), "",
          "_Pemeriksaan render visual (teks meluap, kontras) tidak dijalankan di aplikasi: buka .pptx dan periksa tiap slide._"]
    s.pres_brief, s.pres_fakta, s.pres_qa = brief, fakta, "\n".join(qa)
    s.pres_plan, s.pres_pptx, s.pres_warn, s.pres_template = json.dumps(plan, ensure_ascii=False, indent=1), pptx_bytes, warns, template
    bump("pres_plan")
    return True


def rebuild_presentation():
    s = st.session_state
    try:
        plan = parse_plan(s.pres_plan)
        s.pres_pptx, s.pres_warn = build_pptx(plan, s.figures, s.pres_template)
    except Exception as e:  # noqa: BLE001
        st.error(f"Gagal membangun ulang .pptx: {e}")


for _k, _v in {"pres_skill": "", "pres_brief": "", "pres_fakta": "", "pres_plan": "", "pres_qa": "",
               "pres_pptx": b"", "pres_warn": [], "pres_template": None}.items():
    st.session_state.setdefault(_k, _v)

with tab5:
    sec("🎤", "Buat presentasi", "Bahan presentasi (.pptx) dari naskah final di tab 4, memakai skill presentasi-dari-dokumen.")
    if not any(s.sections.values()):
        empty_state("🎤", "Belum ada naskah", "Selesaikan tab 3 (Tulis) dan tab 4 (Review); presentasi dibuat dari naskah final.")
    else:
        p1, p2, p3 = st.columns([2, 1, 1])
        pres_aud = p1.text_input("Audiens", value="Komunitas ilmiah dan pemangku kepentingan", key="pres_aud")
        pres_goal = p2.selectbox("Tujuan", ["Informasi", "Keputusan", "Persetujuan"], key="pres_goal")
        pres_dur = p3.number_input("Durasi (menit)", min_value=5, max_value=120, value=15, step=5, key="pres_dur")
        pres_refs = st.file_uploader(
            "Dokumen rujukan (opsional): template .pptx, pedoman format, KAK, contoh deck",
            type=FILE_TYPES + ["pptx"], accept_multiple_files=True, key="pres_refs",
        )
        st.caption("Isi slide hanya dari naskah final. Rujukan mengatur bentuk; .pptx pertama dipakai sebagai template.")
        if st.button("🎤 Buat presentasi", type="primary", use_container_width=True):
            if not s.api_key:
                st.error("Isi API key di sidebar terlebih dahulu.")
            else:
                generate_presentation(figs_to_text(assemble_manuscript()), pres_refs, pres_aud, pres_goal, int(pres_dur))

    if s.pres_pptx:
        sec("📦", "Hasil presentasi", "Unduh, lalu periksa tiap slide sebelum dipakai.")
        st.download_button(
            "⬇️ Unduh .pptx", data=s.pres_pptx, file_name="presentasi.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", use_container_width=True,
        )
        for _w in s.pres_warn:
            st.warning(_w)
        with st.expander("👀 Pratinjau isi slide", expanded=True):
            try:
                for _i, _sl in enumerate(parse_plan(s.pres_plan)["slides"], 1):
                    st.markdown(f"**{_i}. {_sl.get('judul', '')}** · _{_sl.get('tipe', 'isi')}_")
                    for _p in _sl.get("poin") or []:
                        st.markdown(f"- {_p}")
            except Exception:  # noqa: BLE001
                st.info("Peta slide belum valid; perbaiki JSON di bawah lalu bangun ulang.")
        with st.expander("🗺️ Peta slide (JSON, bisa diedit)"):
            editor("Peta slide", "pres_plan", 360)
            if st.button("🔄 Bangun ulang .pptx dari peta slide"):
                rebuild_presentation()
                st.rerun()
        with st.expander("🧪 Laporan QA"):
            st.markdown(s.pres_qa)
        with st.expander("📚 Brief rujukan"):
            st.markdown(s.pres_brief)
        with st.expander("🧾 Fakta kunci"):
            st.markdown(s.pres_fakta)


# Isi indikator tahap di akhir agar mencerminkan state terbaru pada run ini
steps_ph.markdown(steps_html(), unsafe_allow_html=True)
